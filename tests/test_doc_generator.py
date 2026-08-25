from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from docx import Document

from doc_generator import BidDocumentGenerator


def test_generator_paths_are_anchored_to_base(tmp_path):
    generator = BidDocumentGenerator(tmp_path)
    assert generator.base_path == Path(tmp_path)
    assert (tmp_path / "templates").is_dir()
    assert (tmp_path / "output").is_dir()
    assert (tmp_path / "assets").is_dir()


def test_partner_count_rejects_missing_middle_partner(tmp_path):
    generator = BidDocumentGenerator(tmp_path)
    with pytest.raises(ValueError, match="first partner"):
        generator.determine_partner_count(
            {
                "LEAD_PARTNER_NAME": "Lead Co",
                "FIRST_PARTNER_NAME": "",
                "SECOND_PARTNER_NAME": "Second Co",
            }
        )


def test_partner_count_supports_one_two_and_three_partners(tmp_path):
    generator = BidDocumentGenerator(tmp_path)
    assert generator.determine_partner_count({"LEAD_PARTNER_NAME": "Lead"}) == 1
    assert generator.determine_partner_count(
        {"LEAD_PARTNER_NAME": "Lead", "FIRST_PARTNER_NAME": "First"}
    ) == 2
    assert generator.determine_partner_count(
        {
            "LEAD_PARTNER_NAME": "Lead",
            "FIRST_PARTNER_NAME": "First",
            "SECOND_PARTNER_NAME": "Second",
        }
    ) == 3


def test_unresolved_placeholder_detection(tmp_path):
    generator = BidDocumentGenerator(tmp_path)
    document = Document()
    document.add_paragraph("Hello {{KNOWN}} and {{MISSING}}")
    generator.replace_in_document(document, {"{{KNOWN}}": "World"})
    assert generator.unresolved_placeholders(document) == ["{{MISSING}}"]
